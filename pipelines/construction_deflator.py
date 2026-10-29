"""建設工事費デフレーター（月次・年度次）の取得。

国土交通省 総合政策局 建設経済統計調査室が毎月公表する建設工事費デフレーターの Excel を取得し、
工事種類が横に並ぶ表を縦持ちの Parquet にする。

- 月次      2016年4月以降。毎月末に2か月前の分が足される
- 年度次    1951年度以降。古い年度ほど公表されている工事種類が少なく、無い値は「***」

列の見出しは8段で、上の段ほど大きな括り（建設総合 > 建築総合 > 住宅総合 > ...）になる。
括りが途中で終わる列は、残りの段に同じ名前が繰り返し入っている。
見出しの右端にある W在来住宅 などの15列は、上の括りを持たない構造別・用途別の系列。

過去の月は Excel 上で行が非表示（グループ化）になっているだけで、値はシートの中にある。

出典: 「建設工事費デフレーター」（国土交通省）
https://www.mlit.go.jp/statistics/details/t-other-2_tk_000362.html
"""

import logging
import re
import unicodedata
from pathlib import Path

import duckdb
import openpyxl

from shared.fetch import fetch, replacing, write_atomic

logger = logging.getLogger("pipelines")

PAGE_URL = "https://www.mlit.go.jp/statistics/details/t-other-2_tk_000362.html"
BASE_URL = "https://www.mlit.go.jp"

# ファイル名は公表日で変わる（deftsuki_2607.xlsx、defnendo_260630.xlsx など）ので、接頭辞で引く
PREFIXES = {
    "monthly": "deftsuki_",
    "fiscal_year": "defnendo_",
}

DATA_DIR = Path("data")
XLSX_DIR = DATA_DIR / "xlsx"
PARQUET_DIR = DATA_DIR / "parquet"

USER_AGENT = "queria-dataset-mlit (+https://queria.io/)"
REQUEST_TIMEOUT = 60

# 見出しの段（0 始まりの行番号）
HEADER_ROWS = range(2, 10)

COLUMNS = {
    "category_order": "INTEGER",
    "category": "VARCHAR",
    "category_path": "VARCHAR",
    "parent_category": "VARCHAR",
    "level": "INTEGER",
    "year": "INTEGER",
    "month": "INTEGER",
    "fiscal_year": "INTEGER",
    "is_provisional": "BOOLEAN",
    "deflator": "DOUBLE",
}


def download_construction_deflator() -> None:
    """建設工事費デフレーターの Excel（月次・年度次）をダウンロードする。

    毎月の公表で中身が入れ替わるので、取得済みでも取り直す。
    """
    page = fetch(PAGE_URL, user_agent=USER_AGENT, timeout=REQUEST_TIMEOUT).decode("utf-8")
    XLSX_DIR.mkdir(parents=True, exist_ok=True)
    for name, prefix in PREFIXES.items():
        match = re.search(r'href="([^"]*/' + re.escape(prefix) + r'[^"/]*\.xlsx)"', page)
        if match is None:
            raise RuntimeError(f"{PAGE_URL}: {prefix}*.xlsx が見つからない")
        href = match.group(1)
        url = href if href.startswith("http") else BASE_URL + href
        body = fetch(url, user_agent=USER_AGENT, timeout=REQUEST_TIMEOUT)
        write_atomic(XLSX_DIR / f"construction_deflator_{name}.xlsx", body)
        logger.info("construction deflator: downloaded %s (%s bytes)", url, len(body))


def _normalize(text: str) -> str:
    """半角カナ・全角英数をそろえ、改行と空白を落とす。"""
    return re.sub(r"\s+", "", unicodedata.normalize("NFKC", text))


def _categories(rows: list[tuple], first: int) -> list[tuple[int, list[str]]]:
    """見出しの8段から、列ごとの工事種類の括り（上から順の名前の並び）を作る。

    first は値が始まる列。返すのは (列, 括り) の並び。
    """
    categories = []
    seen = set()
    for column in range(first, len(rows[HEADER_ROWS[0]])):
        names = [rows[row][column] for row in HEADER_ROWS]
        if all(name is None for name in names):
            continue
        path = []
        for name in names:
            if name is None:
                continue
            name = _normalize(str(name))
            if not path or path[-1] != name:
                path.append(name)
        # その他土木の内訳の「その他土木」のように、括りと同じ名前の内訳は段を畳むと
        # 括りそのものの列と区別できなくなる。先に出た列と重なったら、内訳として1段足す
        if tuple(path) in seen:
            path.append(path[-1])
        if tuple(path) in seen:
            raise ValueError(f"列 {column} の工事種類が他の列と区別できない: {path}")
        seen.add(tuple(path))
        categories.append((column, path))
    return categories


def _value(cell) -> float | None:
    if isinstance(cell, (int, float)):
        return float(cell)
    if cell is None or str(cell).strip() == "***":
        return None
    raise ValueError(f"想定しない値: {cell!r}")


def _fiscal_year(cell) -> tuple[int, bool] | None:
    """年度の列を読む。2017 までは数値、2018 からは「2018年度」「2023年度（暫定）」の文字列。"""
    if isinstance(cell, int):
        return cell, False
    if isinstance(cell, str):
        match = re.fullmatch(r"(\d{4})年度(（暫定）)?", cell.strip())
        if match:
            return int(match.group(1)), match.group(2) is not None
    return None


def _records(categories, values, year, month, fiscal_year, is_provisional) -> list[dict]:
    records = []
    for order, (column, path) in enumerate(categories, start=1):
        value = _value(values[column])
        if value is None:
            continue
        records.append(
            {
                "category_order": order,
                "category": path[-1],
                "category_path": " > ".join(path),
                "parent_category": path[-2] if len(path) > 1 else None,
                "level": len(path),
                "year": year,
                "month": month,
                "fiscal_year": fiscal_year,
                "is_provisional": is_provisional,
                "deflator": value,
            }
        )
    return records


def _parse_monthly(path: Path) -> list[dict]:
    workbook = openpyxl.load_workbook(path, read_only=True, data_only=True)
    sheet = next(sheet for sheet in workbook.worksheets if "月別" in sheet.title)
    rows = list(sheet.iter_rows(values_only=True))
    if rows[HEADER_ROWS[0]][:2] != ("年", "月"):
        raise ValueError(f"{path}: 月別シートの見出しが想定と違う: {rows[HEADER_ROWS[0]][:2]!r}")
    categories = _categories(rows, first=2)
    records = []
    for row in rows[HEADER_ROWS[-1] + 1 :]:
        year, month = row[0], row[1]
        if not (isinstance(year, int) and isinstance(month, int)):
            continue
        fiscal_year = year if month >= 4 else year - 1
        # 年度末までの月は行だけ先に用意されていて、値が空
        records += _records(categories, row, year, month, fiscal_year, None)
    return records


def _parse_fiscal_year(path: Path) -> list[dict]:
    workbook = openpyxl.load_workbook(path, read_only=True, data_only=True)
    sheet = workbook.worksheets[0]
    rows = list(sheet.iter_rows(values_only=True))
    if rows[HEADER_ROWS[0]][0] != "年度":
        raise ValueError(f"{path}: 年度別シートの見出しが想定と違う: {rows[HEADER_ROWS[0]][0]!r}")
    categories = _categories(rows, first=1)
    records = []
    for row in rows[HEADER_ROWS[-1] + 1 :]:
        parsed = _fiscal_year(row[0])
        if parsed is None:
            continue
        fiscal_year, is_provisional = parsed
        records += _records(categories, row, None, None, fiscal_year, is_provisional)
    return records


PARSERS = {
    "monthly": _parse_monthly,
    "fiscal_year": _parse_fiscal_year,
}


def build_construction_deflator_parquet() -> None:
    """Excel を縦持ちにして Parquet に書く。1行 = 工事種類 × 期間。"""
    PARQUET_DIR.mkdir(parents=True, exist_ok=True)
    for name, parse in PARSERS.items():
        records = parse(XLSX_DIR / f"construction_deflator_{name}.xlsx")
        destination = PARQUET_DIR / f"construction_deflator_{name}.parquet"
        with replacing(destination) as partial:
            if not records:
                raise RuntimeError(f"construction deflator ({name}): 0 rows")
            connection = duckdb.connect()
            columns = ", ".join(f"{column} {kind}" for column, kind in COLUMNS.items())
            connection.execute(f"CREATE TABLE deflator ({columns})")
            connection.executemany(
                f"INSERT INTO deflator VALUES ({', '.join('?' for _ in COLUMNS)})",
                [tuple(record[column] for column in COLUMNS) for record in records],
            )
            connection.execute(
                "COPY deflator TO $destination (FORMAT parquet)",
                {"destination": str(partial)},
            )
        logger.info("construction deflator (%s): %d rows -> %s", name, len(records), destination)
